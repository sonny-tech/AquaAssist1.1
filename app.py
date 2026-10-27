import os
import re
import math
import streamlit as st
import pdfplumber
from docx import Document
from fpdf import FPDF

import ui

# ------------------ CONFIG ------------------

ui.setup_page(st.session_state.get("step", "program"))  # page config + styling; must come first

LEVELS_DIR = "skill_tracking"       # PDFs named after levels e.g. "Tiny Tots Level 5 Sting Rays Must Sees.pdf"
PROGRESSIONS_DIR = "lesson_progressions"  # Docx files named after skills e.g. "Back Crawl.docx"
OUTPUT_DIR = "output"
FONT_DIR = "fonts"  # put DejaVuSans.ttf and DejaVuSans-Bold.ttf here
os.makedirs(OUTPUT_DIR, exist_ok=True)


# ------------------ FILE DISCOVERY ------------------

def get_all_progression_names():
    """Return list of skill names from docx filenames (without extension)."""
    if not os.path.exists(PROGRESSIONS_DIR):
        return []
    return [
        os.path.splitext(f)[0]
        for f in os.listdir(PROGRESSIONS_DIR)
        if f.lower().endswith(".docx")
    ]

# Order the level groups appear in on screen; any other groups follow alphabetically.
GROUP_ORDER = ["Tiny Tots", "Swimmer", "Junior Masters"]


def list_levels():
    """Every level PDF in LEVELS_DIR as {"name", "path", "group", "number"}, sorted by number.

    'Tiny_Tots_Level_5_Sting_Rays_Must_Sees.pdf' -> name 'Tiny Tots Level 5 Sting Rays',
    group 'Tiny Tots'. 'Swimmer_3_Must_Sees.pdf' -> name 'Swimmer 3', group 'Swimmer'.
    """
    levels = []
    if not os.path.isdir(LEVELS_DIR):
        return levels
    for f in os.listdir(LEVELS_DIR):
        if not f.lower().endswith(".pdf"):
            continue
        stem = os.path.splitext(f)[0]
        name = re.sub(r"[_\s]+must[_\s]*sees$", "", stem, flags=re.IGNORECASE)
        name = " ".join(name.replace("_", " ").split())
        digit = re.search(r"\d+", name)
        prefix = name[: digit.start()].strip() if digit else name
        group = re.sub(r"\s+level$", "", prefix, flags=re.IGNORECASE).strip() or name
        levels.append({
            "name": name,
            "path": os.path.join(LEVELS_DIR, f),
            "group": group,
            "number": int(digit.group()) if digit else 0,
        })
    levels.sort(key=lambda l: (l["number"], l["name"].lower()))
    return levels


def ordered_groups(levels):
    groups = sorted({l["group"] for l in levels}, key=str.lower)
    preferred = [g for want in GROUP_ORDER for g in groups if g.lower() == want.lower()]
    return preferred + [g for g in groups if g not in preferred]


def skills_for_level(level_path):
    """The level's must-sees as pickable skills, grouped under the PDF's headings, plus every other skill.

    A must-see that starts with the name of a skill that has a progression doc is listed under that
    skill's name (synonyms count, so 'Tread water 30 sec.' finds Treading Water). A must-see with no
    progression doc is listed as itself, and goes into the plan as written.
    Returns (sections, others): sections is [(heading, [skill, ...])] in PDF order, and others
    are the progression-doc skills this level's must-sees don't mention.
    """
    known = sorted(get_all_progression_names(), key=str.lower)
    keys = {s: normalize_name(s) for s in known}
    sections, seen = [], set()
    for heading, text in parse_must_sees(level_path):
        line = normalize_name(text)
        # The line has to start with the skill's name, so 'Front Crawl 15m' is the Front Crawl skill but
        # 'Interval training: 4 x 15m flutter kick...' stays its own must-see.
        hits = [s for s in known if keys[s] and line.startswith(keys[s])]
        # If two names both fit, the longer one wins (Legs Only Treading Water over a shorter name).
        hits = [s for s in hits if not any(keys[s] != keys[o] and keys[o].startswith(keys[s]) for o in hits)]
        for skill in hits or [text]:
            if skill in seen:
                continue
            seen.add(skill)
            if not sections or sections[-1][0] != heading:
                sections.append((heading, []))
            sections[-1][1].append(skill)
    others = [s for s in known if s not in seen]
    return sections, others


# A bullet at the start of a line: a real bullet, the "(cid:127)" some PDFs extract it as, or a dash.
_BULLET = re.compile(r"^\s*(?:•|\(cid:\d+\)|[·▪●◦■‣*]|[-–—])\s*")


def parse_must_sees(pdf_path):
    """Read a level's must-sees PDF into [(heading, text)], one entry per bulleted line.

    Lines without a bullet are headings ('Surface Support') and become the heading of the bullets
    that follow; the first line is the level's title and is skipped. A line that starts in lower
    case right after a bullet is the wrapped end of that bullet.
    """
    rows = []  # ["item" | "head", text]
    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            for raw in (page.extract_text() or "").split("\n"):
                raw = raw.strip()
                if not raw:
                    continue
                bullet = _BULLET.match(raw)
                if bullet:
                    rows.append(["item", raw[bullet.end():].strip()])
                elif rows and rows[-1][0] == "item" and raw[0].islower():
                    rows[-1][1] += " " + raw
                else:
                    rows.append(["head", raw])
    if not any(kind == "item" for kind, _ in rows):  # no bullets at all: every line is a must-see
        return [("", text) for _, text in rows]
    items, heading = [], ""
    for i, (kind, text) in enumerate(rows):
        if kind == "head":
            if i > 0:
                heading = text.rstrip(":").strip()
        elif text:
            items.append((heading, text))
    return items


def get_must_sees_from_level(pdf_path):
    """Just the text of each must-see line (used to look up distances and times)."""
    return [text for _, text in parse_must_sees(pdf_path)]


# Alternate names swimmers/parents commonly use that should be treated as
# the same skill as the canonical docx/progression name. Add more pairs
# here as needed (alias -> canonical name, case-insensitive).
# Alternate names swimmers/parents commonly use that should be treated as
# the same skill as the canonical docx/progression name. Add more pairs
# here as needed (alias -> canonical name, case-insensitive).
SKILL_SYNONYMS = {
    "freestyle": "front crawl",
    "backstroke": "back crawl",
    "back paddle": "back kick w scull/paddle",
    "back kick with scull/paddle": "back kick w scull/paddle",
    "front glide": "glide on front",
    "back glide": "glide on back",
    "front float": "float on front",
    "back float": "float on back",
    "flutter kick": "flutter kick on front",
    "back kick": "flutter kick on back",
    "side glide": "flutter kick on side",
    "treading water": "tread water",
    "legs-only surface support": "legs only tread water",
    "legs only surface support": "legs only tread water",
    "jump entry into deep water and tread": "jump entry into deep water and tread water",
}

# Some short aliases above are also literal prefixes of *other* distinct
# skill names (e.g. "back kick" is a prefix of "back kick w scull/paddle",
# and "flutter kick" is a prefix of "flutter kick on back/side"). Without
# an exclusion, substituting the short alias would corrupt those longer,
# unrelated names. Each entry here is a negative-lookahead suffix appended
# to that alias's match pattern so it only fires when NOT followed by the
# start of one of those longer names.
_SYNONYM_EXCLUDE_AFTER = {
    "back kick": r"(?!\s*(?:w|with)\W*scull\W*paddle)",
    "flutter kick": r"(?!\s*on\s*(?:back|front|side))",
    "jump entry into deep water and tread": r"(?!\s*water)",
}

def canonicalize_skill_text(text):
    """Replace any whole-word synonym occurrence with its canonical name,
    e.g. 'Freestyle 15m' -> 'front crawl 15m'. Case-insensitive; operates
    on lowercased text with word boundaries. Longest aliases are applied
    first so multi-word aliases win over any shorter alias that happens to
    be one of their prefixes, and _SYNONYM_EXCLUDE_AFTER further protects
    short aliases from matching inside unrelated longer skill names."""
    result = text.lower()
    for alias in sorted(SKILL_SYNONYMS, key=len, reverse=True):
        canonical = SKILL_SYNONYMS[alias]
        exclude = _SYNONYM_EXCLUDE_AFTER.get(alias, "")
        pattern = r"\b" + re.escape(alias) + r"\b" + exclude
        result = re.sub(pattern, canonical, result)
    return result


def normalize_name(text):
    """Canonicalize synonyms, then lowercase and strip all non-alphanumeric
    characters (including spaces) for loose skill-name matching — makes
    'Front Crawl', 'FrontCrawl', and 'Freestyle' all resolve to the same
    key ('frontcrawl')."""
    return re.sub(r"[^a-z0-9]", "", canonicalize_skill_text(text))


# Matches the first number+unit occurring anywhere in a line, e.g.:
#   "Front crawl 15m"                        -> "15"
#   "Front Crawl 5m with bilateral breathing" -> "5"   (extra text after)
#   "Unassisted rocketship (5m)"              -> "5"   (parentheses)
#   "Tread water 30 sec."                     -> "30"
_NUMBER_UNIT_PATTERN = re.compile(
    r"(\d+(?:\.\d+)?)\s*(m|sec\.?|secs\.?|seconds?|min\.?|mins\.?|minutes?)\b",
    re.IGNORECASE,
)

def _value_from_match(match):
    """Return the number as a string, converting minutes to seconds
    (e.g. '1 min' -> '60') since time blanks in the lesson plan are in
    seconds. Meters and seconds are returned as-is."""
    number, unit = match.group(1), match.group(2).lower()
    if unit.startswith("min"):
        seconds = float(number) * 60
        return str(int(seconds)) if seconds == int(seconds) else f"{seconds:g}"
    return number

def _unit_kind(unit):
    """'time' for seconds/minutes, 'distance' for meters."""
    return "time" if unit.lower().startswith(("sec", "min")) else "distance"

# Finds the unit written right after a '_' blank in a progression line,
# e.g. 'Shoulder roll w KB _ m (2 min)' -> 'm', 'Front Crawl _ x 2 m' -> 'm',
# 'Tread water _ sec' -> 'sec'. Stops at '(' so the '(X min)' time tag is ignored.
_BLANK_UNIT_PATTERN = re.compile(
    r"_[^()]*?\b(m|sec\.?|secs\.?|seconds?|min\.?|mins\.?|minutes?)\b",
    re.IGNORECASE,
)

def blank_kind(line):
    """What kind of value a '_' blank in this line expects: 'distance'
    (meters), 'time' (seconds/minutes), or None if no unit follows it."""
    match = _BLANK_UNIT_PATTERN.search(line)
    return _unit_kind(match.group(1)) if match else None


def find_distance_for_skill(skill_name, must_sees_lines, kind=None):
    """Find the must-sees line mentioning this skill and pull its first
    number+unit, wherever it falls in the line (not just at the end).

    Real must-sees PDFs are inconsistent — some lines end right after the
    distance, some have trailing descriptive text, some wrap the distance
    in parentheses. Matching on 'contains the skill name' + 'first number
    found in that line' handles all of those instead of requiring an exact
    line format.

    If kind is 'distance' or 'time', only numbers with that kind of unit
    are considered, so a blank that expects meters never picks up a
    seconds value (and vice versa). kind=None accepts either.
    """
    target = normalize_name(skill_name)
    if not target:
        return None

    normalized_lines = [(line, normalize_name(line)) for line in must_sees_lines]

    # Prefer lines that START with the skill name (so 'Tread Water' picks
    # 'Tread water 30 sec.' rather than a longer line that merely contains
    # it, like 'Legs-only surface support'); fall back to 'contains'.
    for matches_line in (
        lambda norm: norm.startswith(target),
        lambda norm: target in norm,
    ):
        for line, norm in normalized_lines:
            if matches_line(norm):
                for match in _NUMBER_UNIT_PATTERN.finditer(line):
                    if kind is None or _unit_kind(match.group(2)) == kind:
                        return _value_from_match(match)
    return None


# ------------------ PROGRESSION PARSING ------------------

def get_progressions_for_skill(skill_name):
    """Read progressions list from the skill's docx file."""
    path = os.path.join(PROGRESSIONS_DIR, f"{skill_name}.docx")
    if not os.path.exists(path):
        return []
    doc = Document(path)
    progressions = []
    for para in doc.paragraphs:
        text = para.text.strip().lstrip("-").strip()
        if text:
            progressions.append(text)
    return progressions


# ------------------ EQUIPMENT DETECTION ------------------

# Maps a lowercase keyword (matched as a whole word) to its display label.
# "kb" is the common shorthand for kickboard seen in progression text.
EQUIPMENT_KEYWORDS = {
    "marshmallow": "Marshmallow",
    "marshmallows": "Marshmallow",
    "kickboard": "Kickboard",
    "kickboards": "Kickboard",
    "kb": "Kickboard",
    "noodle": "Noodle",
    "noodles": "Noodle",
    "mat": "Mat",
    "mats": "Mat",
}

def detect_equipment(skills_and_progressions):
    """Scan skill names + progression text for equipment keywords.

    Uses whole-word matching (not substring) so e.g. 'mat' doesn't
    false-positive inside words like 'format' or 'automatic'.
    """
    combined_text = " ".join(
        skill + " " + " ".join(progressions)
        for skill, progressions in skills_and_progressions.items()
    ).lower()

    found = set()
    for keyword, label in EQUIPMENT_KEYWORDS.items():
        if re.search(r"\b" + re.escape(keyword) + r"\b", combined_text):
            found.add(label)

    return sorted(found)


# ------------------ TIME CALCULATION ------------------

_TIME_TAG_PATTERN = re.compile(r"\((\d+(?:\.\d+)?)\s*min\)", re.IGNORECASE)

_SECONDS_PATTERN = re.compile(r"(\d+(?:\.\d+)?)\s*(?:sec\.?|secs\.?|seconds?)\b", re.IGNORECASE)

def calculate_total_minutes(skills_and_progressions):
    """Total lesson minutes: every '(X min)' time tag, plus the seconds-based
    progressions (e.g. treading water) — those seconds are added together,
    divided by 60, and rounded up. Lines that already carry a '(X min)' tag
    are counted by the tag only, so nothing is double counted."""
    total_minutes = 0.0
    total_seconds = 0.0
    for progressions in skills_and_progressions.values():
        for line in progressions:
            tag = _TIME_TAG_PATTERN.search(line)
            if tag:
                total_minutes += float(tag.group(1))
                continue
            seconds = _SECONDS_PATTERN.search(line)
            if seconds:
                total_seconds += float(seconds.group(1))
    return total_minutes + math.ceil(total_seconds / 60)


# ------------------ PDF GENERATION ------------------

class LessonPlanPDF(FPDF):
    def __init__(self):
        super().__init__()
        # Register Unicode fonts so em dashes, smart quotes, accents, etc.
        # never blow up FPDF's default Latin-1 encoder.
        regular = os.path.join(FONT_DIR, "DejaVuSans.ttf")
        bold = os.path.join(FONT_DIR, "DejaVuSans-Bold.ttf")
        italic = os.path.join(FONT_DIR, "DejaVuSans-Oblique.ttf")

        self.add_font("DejaVu", "", regular)
        self.add_font("DejaVu", "B", bold if os.path.exists(bold) else regular)
        self.add_font("DejaVu", "I", italic if os.path.exists(italic) else regular)

    def header(self):
        start_y = self.get_y()

        # Total lesson time, top-right corner
        if self.total_time:
            self.set_font("DejaVu", "", 10)
            self.set_text_color(90, 90, 90)
            self.set_xy(-70, start_y)
            self.cell(60, 6, f"Total Time: {self.total_time}", align="R")

        self.set_xy(10, start_y)
        self.set_font("DejaVu", "B", 14)
        self.set_text_color(31, 78, 121)  # Navy blue
        self.cell(0, 10, self.title, ln=True, align="C")
        self.set_draw_color(31, 78, 121)
        self.set_line_width(0.5)
        self.line(10, self.get_y(), 200, self.get_y())
        self.ln(4)

    def footer(self):
        self.set_y(-15)
        self.set_font("DejaVu", "I", 8)
        self.set_text_color(150, 150, 150)
        self.cell(0, 10, f"AquaAssist 2.0 — Page {self.page_no()}", align="C")

def generate_lesson_plan_pdf(level_name, skills_and_progressions, output_path):
    pdf = LessonPDF()
    pdf.title = f"{level_name} — Lesson Plan"

    total_minutes = calculate_total_minutes(skills_and_progressions)
    pdf.total_time = (
        f"{int(total_minutes)} min" if total_minutes == int(total_minutes) else f"{total_minutes:g} min"
    ) if total_minutes > 0 else None

    pdf.add_page()

    # Equipment Needed section
    equipment = detect_equipment(skills_and_progressions)
    if equipment:
        pdf.set_font("DejaVu", "B", 12)
        pdf.set_fill_color(214, 228, 240)
        pdf.set_text_color(31, 78, 121)
        pdf.cell(0, 8, "Equipment Needed", ln=True, fill=True)
        pdf.ln(1)

        pdf.set_font("DejaVu", "", 10)
        pdf.set_text_color(26, 26, 26)
        pdf.cell(8)
        pdf.cell(0, 7, ", ".join(equipment), ln=True)
        pdf.ln(3)

    for skill, progressions in skills_and_progressions.items():
        # Skill header
        pdf.set_font("DejaVu", "B", 12)
        pdf.set_fill_color(214, 228, 240)
        pdf.set_text_color(31, 78, 121)
        pdf.cell(0, 8, skill, ln=True, fill=True)
        pdf.ln(1)

        # Progressions
        pdf.set_font("DejaVu", "", 10)
        pdf.set_text_color(26, 26, 26)
        if progressions:
            for i, prog in enumerate(progressions, start=1):
                pdf.cell(8)  # indent
                pdf.cell(0, 7, f"{i}. {prog}", ln=True)
        else:
            pdf.cell(8)
            pdf.set_text_color(150, 150, 150)
            pdf.cell(0, 7, "No progressions found for this skill.", ln=True)
            pdf.set_text_color(26, 26, 26)

        pdf.ln(3)

    pdf.output(output_path)


# Monkey-patch title/total_time onto FPDF since we set them manually
class LessonPDF(LessonPlanPDF):
    title = ""
    total_time = None


# ------------------ LESSON PLAN GENERATION ------------------

OTHER_SECTION = "Other must-sees"  # heading for must-sees that have no progression doc yet
DEFAULT_MINUTES = 5               # time given to a must-see from the sheet that has no time of its own


def format_minutes(total_minutes):
    if total_minutes == int(total_minutes):
        return f"{int(total_minutes)} min"
    return f"{total_minutes:g} min"


def build_lesson_plan(level_name, level_path, chosen_skills):
    """Run the whole pipeline for the chosen level and skills and return what to show.

    Nothing is drawn here. The result is a dict that render_plan() displays,
    and that "Recent plans" can reopen later:
      notes:      [(kind, text)] where kind is error / warning / info / success
      skills:     the skills included in the plan
      total_time: e.g. "12 min", or None
      pdf:        the lesson plan PDF as bytes, or None if nothing was generated
    """
    entry = {
        "level_name": level_name,
        "skills": [],
        "notes": [],
        "total_time": None,
        "pdf": None,
        "file_name": None,
    }
    notes = entry["notes"]

    # Pass 1: collect each skill's progressions and whatever values the level's
    # must-sees list gives it, without filling blanks yet. Distances (meters)
    # and times (seconds) are looked up separately so a blank expecting meters
    # never gets a seconds value, or the reverse.
    must_sees_lines = get_must_sees_from_level(level_path)
    known = set(get_all_progression_names())
    docx_skills = [s for s in chosen_skills if s in known]
    as_written = [s for s in chosen_skills if s not in known]  # must-sees with no progression doc
    raw_progressions = {}
    skill_values = {}  # {skill: {"distance": .., "time": .., None: ..}}
    for skill in docx_skills:
        raw_progressions[skill] = get_progressions_for_skill(skill)
        skill_values[skill] = {
            kind: find_distance_for_skill(skill, must_sees_lines, kind)
            for kind in ("distance", "time", None)
        }

    # Flag a genuine conflict if both Front Crawl and Back Crawl were chosen
    # with different explicit distances. Don't silently override, just say so.
    front_key = next(
        (s for s in raw_progressions if normalize_name(s) == normalize_name("Front Crawl")), None
    )
    back_key = next(
        (s for s in raw_progressions if normalize_name(s) == normalize_name("Back Crawl")), None
    )
    if front_key and back_key:
        front_dist = skill_values[front_key]["distance"]
        back_dist = skill_values[back_key]["distance"]
        if front_dist and back_dist and front_dist != back_dist:
            notes.append((
                "warning",
                f"Front Crawl ({front_dist}m) and Back Crawl ({back_dist}m) have different "
                "distances in the must-sees PDF. Each is used as listed. Worth double-checking.",
            ))

    # Fallback distance (meters only) for any blank that has none of its own:
    # Front Crawl's distance for this level (the usual benchmark), or Flutter
    # Kick's if Front Crawl isn't listed either. Blanks measured in seconds
    # (e.g. treading water) never use these, since they're meters.
    fallback_front_crawl = find_distance_for_skill("Front Crawl", must_sees_lines, "distance")
    fallback_flutter_kick = find_distance_for_skill("Flutter Kick", must_sees_lines, "distance")

    # Pass 2: fill in each blank according to the unit written after it.
    filled_skills = {}
    missing_distance = []
    used_fallback = {}  # {"Front Crawl" | "Flutter Kick": [skill, ...]}

    for skill, progressions in raw_progressions.items():
        filled = []
        unfilled = False
        for p in progressions:
            if "_" not in p:
                filled.append(p)
                continue
            kind = blank_kind(p)
            value = skill_values[skill][kind]
            fallback_source = None
            if not value and kind != "time":
                if fallback_front_crawl:
                    value, fallback_source = fallback_front_crawl, "Front Crawl"
                elif fallback_flutter_kick:
                    value, fallback_source = fallback_flutter_kick, "Flutter Kick"
            if value:
                p = p.replace("_", value)
                if fallback_source and skill not in used_fallback.setdefault(fallback_source, []):
                    used_fallback[fallback_source].append(skill)
            else:
                unfilled = True
            filled.append(p)
        if unfilled:
            missing_distance.append(skill)
        filled_skills[skill] = filled

    for source, skills_list in used_fallback.items():
        fallback_value = fallback_front_crawl if source == "Front Crawl" else fallback_flutter_kick
        notes.append((
            "info",
            f"No distance listed for {', '.join(skills_list)}. Used the {source} "
            f"distance ({fallback_value}m) for this level instead.",
        ))

    if missing_distance:
        notes.append((
            "warning",
            f"No distance found in the must-sees PDF for: {', '.join(missing_distance)}. "
            "Those progression lines still show a blank.",
        ))

    # Must-sees with no progression doc are pasted in as written, in one section at the end.
    # They have no time of their own, so each gets DEFAULT_MINUTES (shown as "(5 min)" like the others).
    plan_sections = dict(filled_skills)
    if as_written:
        plan_sections[OTHER_SECTION] = [
            s if _TIME_TAG_PATTERN.search(s) else f"{s} ({DEFAULT_MINUTES} min)" for s in as_written
        ]

    if not plan_sections:
        notes.append(("error", "No skills were chosen, so there is nothing to put in a plan."))
        return entry

    # Generate the PDF
    safe_level = level_name.replace(" ", "_").replace("/", "-")
    output_path = os.path.join(OUTPUT_DIR, f"{safe_level}_lesson_plan.pdf")
    generate_lesson_plan_pdf(level_name, plan_sections, output_path)

    notes.append(("success", f"Lesson plan ready for {level_name}."))

    total_minutes = calculate_total_minutes(plan_sections)
    entry["total_time"] = format_minutes(total_minutes) if total_minutes > 0 else None
    entry["skills"] = list(chosen_skills)
    entry["file_name"] = os.path.basename(output_path)
    with open(output_path, "rb") as f:
        entry["pdf"] = f.read()
    return entry


# ------------------ STEPS ------------------
# The person moves through four steps: "program" -> "level" -> "skills" -> "plan".
# Button clicks run these callbacks first, then the script reruns and draws the new step.

state = st.session_state
state.setdefault("step", "program")
state.setdefault("program", None)  # e.g. "Swimmer"
state.setdefault("level", None)    # {"name": ..., "path": ...}
state.setdefault("picked", [])     # skills chosen, in the order they were tapped
state.setdefault("entry", None)    # the generated plan for the plan step
state.setdefault("history", [])    # plans made in this browser session, newest first
state.setdefault("next_id", 1)


def choose_program(program):
    state.update(step="level", program=program, level=None, picked=[], entry=None)


def choose_level(name, path):
    state.update(step="skills", level={"name": name, "path": path}, picked=[], entry=None)


def toggle_skill(skill):
    if skill in state.picked:
        state.picked.remove(skill)
    else:
        state.picked.append(skill)


def start_plan():
    state.update(step="plan", entry=None)


def go_program():
    state.update(step="program", program=None, level=None, picked=[], entry=None)


def go_level():
    state.update(step="level", level=None, picked=[], entry=None)


def go_skills():
    state.update(step="skills", entry=None)


def reopen_plan(entry_id):
    entry = next((e for e in state.history if e["id"] == entry_id), None)
    if entry:
        state.update(
            step="plan",
            entry=entry,
            program=entry["program"],
            level={"name": entry["level_name"], "path": entry["level_path"]},
            picked=list(entry["skills"]),
        )


def render_program_step():
    levels = list_levels()
    with st.container(key="aa-panel-program"):
        ui.heading("Choose a program", "Tap the program you are teaching today.")
        if not levels:
            st.warning(f"No level files found in the {LEVELS_DIR} folder.")
        with st.container(key="aa-grid-prog"):
            for i, group in enumerate(ordered_groups(levels)):
                st.button(group, key=f"aa-prog-{i}", on_click=choose_program, args=(group,))

        if state.history:
            ui.label("Recent plans")
            with st.container(key="aa-grid-recent"):
                for entry in state.history[:6]:
                    text = f"{entry['level_name']}: {', '.join(entry['skills'])}"
                    st.button(
                        text if len(text) <= 44 else text[:43].rstrip() + "…",
                        key=f"aa-recent-{entry['id']}",
                        help=text,
                        on_click=reopen_plan,
                        args=(entry["id"],),
                    )


def render_level_step():
    levels = [l for l in list_levels() if l["group"] == state.program]
    numbers = [l["number"] for l in levels]
    plain_numbers = all(numbers) and len(set(numbers)) == len(numbers)
    with st.container(key="aa-panel-level"):
        with st.container(key="aa-actions-top"):
            st.button("Change program", key="aa-nav-program", on_click=go_program)
        ui.heading(f"{state.program}: choose a level", "Tap the level number.")
        with st.container(key="aa-grid-num" if plain_numbers else "aa-grid-names"):
            for i, level in enumerate(levels):
                st.button(
                    str(level["number"]) if plain_numbers else level["name"],
                    key=f"aa-num-{i}",
                    help=level["name"],
                    on_click=choose_level,
                    args=(level["name"], level["path"]),
                )


def skill_grid(skills, offset):
    with st.container(key=f"aa-grid-skills-{offset}"):
        for i, skill in enumerate(skills, start=offset):
            on = skill in state.picked
            st.button(
                skill,
                key=f"aa-skill-{'on' if on else 'off'}-{i}",
                help="Selected. Tap to remove." if on else None,
                on_click=toggle_skill,
                args=(skill,),
            )


def render_skills_step():
    level = state.level
    sections, others = skills_for_level(level["path"])
    count = len(state.picked)
    with st.container(key="aa-panel-skills"):
        with st.container(key="aa-actions-top"):
            st.button("Change level", key="aa-nav-level", on_click=go_level)
        ui.heading(
            f"Choose skills for {level['name']}",
            "Tap as many as you need. They go into the plan in the order you tap them.",
        )
        offset = 0
        for category, skills in sections:
            ui.label(category or "Must-sees")
            skill_grid(skills, offset)
            offset += len(skills)
        if others:
            with st.expander("More skills" if sections else "Skills", expanded=not sections):
                skill_grid(others, offset)
        if not sections and not others:
            st.warning(f"No skills found for this level in the {LEVELS_DIR} or {PROGRESSIONS_DIR} folders.")

        ui.count(f"{count} chosen" if count else "Choose at least one skill")
        st.button("Make my lesson plan", key="aa-go", on_click=start_plan, disabled=not state.picked)


def render_plan_step():
    level = state.level
    entry = state.entry

    if entry is None:
        with ui.loading():
            entry = build_lesson_plan(level["name"], level["path"], state.picked)
        entry["level_path"] = level["path"]
        entry["program"] = state.program
        if entry["pdf"]:
            entry["id"] = state.next_id
            state.next_id += 1
            state.history.insert(0, entry)
        state.entry = entry

    with st.container(key="aa-panel-plan"):
        ui.heading(f"Your lesson plan: {entry['level_name']}")
        for kind, text in entry["notes"]:
            getattr(st, kind)(text)

        stats = []
        if entry["total_time"]:
            value, unit = entry["total_time"].split(" ", 1)
            stats.append((value, unit, "Lesson time"))
        if entry["skills"]:
            n = len(entry["skills"])
            stats.append((str(n), "", "Skills" if n != 1 else "Skill"))
        if stats:
            ui.big_stats(stats)
        if entry["skills"]:
            ui.item_list(entry["skills"])

        with st.container(key="aa-actions-plan"):
            if entry["pdf"]:
                st.download_button(
                    label="Download PDF",
                    data=entry["pdf"],
                    file_name=entry["file_name"],
                    mime="application/pdf",
                    key="aa-download",
                )
            st.button("Change skills", key="aa-nav-skills", on_click=go_skills)
            st.button("New lesson plan", key="aa-nav-new", on_click=go_program)

        if entry["pdf"]:
            ui.pdf_frame(entry["pdf"])


# ------------------ MAIN UI ------------------

# Guard against a half-finished state.
if state.step == "level" and not state.program:
    state.step = "program"
if state.step in ("skills", "plan") and not state.level:
    state.step = "program"

ui.show_splash()
ui.render_header(state.step)

if state.step == "program":
    render_program_step()
elif state.step == "level":
    render_level_step()
elif state.step == "skills":
    render_skills_step()
else:
    render_plan_step()
